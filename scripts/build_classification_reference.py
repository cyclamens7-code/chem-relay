"""
'(참고) 분류표시' 시트(각 물질별 결과물 xlsm의 2번째 시트)를 모아
app/data/classification_reference.json 을 생성한다.

시트 구조 (열 D ~ 시트 마지막 열):
  1행: DB명(+하이퍼링크). 병합된 열은 왼쪽 열에만 값이 있으므로 forward-fill로 채운다.
  B열: 구분(물리적 위험성/건강유해성/환경유해성), C열: 분류표시 항목
  2행~: 각 DB 열에 채워진 값이 그 항목·그 DB의 분류(예: '구분 1B')

CAS·물질명은 '일반정보' 시트에서 찾는다(헤더 행 다음 첫 데이터 행).
같은 CAS가 여러 번 나오면 나중 파일이 앞 파일을 덮어쓴다(최신 제출본 우선 가정).

사용:
  py scripts/build_classification_reference.py
"""
import json
import pathlib
import re
import sys

import openpyxl
from openpyxl.utils import column_index_from_string

ROOT = pathlib.Path(__file__).resolve().parent.parent
SEARCH_DIRS = [
    pathlib.Path(r"C:\Users\user\Desktop\업무용 툴\260714_화학DB사업_DB결과물(6주차)_제출용"
                 r"\260714_화학DB사업_DB결과물_제출용"),
]
OUT_PATH = ROOT / "app" / "data" / "classification_reference.json"

CAS_RE = re.compile(r"^\d{2,7}-\d{2}-\d$")
COL_START = column_index_from_string("D")


def clean_header(v):
    if not v:
        return ""
    s = str(v).replace("\n", " ").replace("[검색]", "").strip()
    s = re.sub(r"\s+", " ", s)
    return s


def find_cas_name(ws):
    """'일반정보' 시트에서 (CAS, 물질명) 을 찾는다. 헤더 'CAS NO.' 다음 데이터 행 사용."""
    header_row = None
    cas_col = name_col = None
    for r in range(1, 20):
        for c in range(1, 10):
            v = ws.cell(row=r, column=c).value
            if v and "CAS" in str(v).upper():
                header_row = r
                cas_col = c
            if v and "화학물질명" in str(v):
                name_col = c
        if header_row:
            break
    if not header_row or not cas_col:
        return None, None
    data_row = None
    for r in range(header_row + 1, header_row + 8):
        v = ws.cell(row=r, column=cas_col).value
        if v and CAS_RE.match(str(v).strip()):
            data_row = r
            break
    if not data_row:
        return None, None
    cas = ws.cell(row=data_row, column=cas_col).value
    name = ws.cell(row=data_row, column=name_col).value if name_col else None
    cas = str(cas).strip() if cas else None
    name = str(name).strip() if name else None
    return cas, name


def extract_file(path: pathlib.Path):
    try:
        wb = openpyxl.load_workbook(path, data_only=True, keep_links=True, read_only=False)
    except Exception as e:
        return None, f"열기 실패: {e}"

    try:
        if "(참고) 분류표시" not in wb.sheetnames or "일반정보" not in wb.sheetnames:
            return None, "필요 시트 없음"
        ws = wb["(참고) 분류표시"]
        ws0 = wb["일반정보"]

        cas, name = find_cas_name(ws0)
        m = CAS_RE.match(path.stem.split("_")[0])
        if not cas or not CAS_RE.match(cas):
            cas = m.group(0) if m else None
        if not cas:
            return None, "CAS 확인 불가"

        # 1행 헤더 forward-fill
        headers = {}
        last_name, last_link = "", None
        for c in range(COL_START, ws.max_column + 1):
            cell = ws.cell(row=1, column=c)
            if cell.value:
                last_name = clean_header(cell.value)
                last_link = cell.hyperlink.target if cell.hyperlink else None
            headers[c] = (last_name, last_link)

        rows = []
        for r in range(2, ws.max_row + 1):
            category = ws.cell(row=r, column=2).value
            item = ws.cell(row=r, column=3).value
            if not item:
                continue
            for c in range(COL_START, ws.max_column + 1):
                val = ws.cell(row=r, column=c).value
                if val in (None, ""):
                    continue
                db_name, db_link = headers.get(c, ("", None))
                if not db_name:
                    continue
                rows.append({
                    "category": str(category).strip() if category else "",
                    "item": str(item).strip(),
                    "grade": str(val).strip(),
                    "db": db_name,
                    "source_url": db_link,
                })
        return {"cas": cas, "name": name, "rows": rows}, None
    finally:
        wb.close()


def main():
    files = []
    for d in SEARCH_DIRS:
        if d.is_dir():
            files.extend(sorted(f for f in d.rglob("*.xlsm") if not f.name.startswith("~$")))
    print(f"대상 파일 {len(files)}개")

    out = {}
    errors = []
    empty_count = 0
    for i, f in enumerate(files, 1):
        data, err = extract_file(f)
        if err:
            errors.append((f.name, err))
            continue
        if not data["rows"]:
            empty_count += 1
        out[data["cas"]] = {"name": data["name"], "rows": data["rows"]}
        if i % 40 == 0:
            print(f"  {i}/{len(files)} 처리...")

    OUT_PATH.parent.mkdir(parents=True, exist_ok=True)
    OUT_PATH.write_text(json.dumps(out, ensure_ascii=False, indent=1), encoding="utf-8")
    print(f"완료: {len(out)}개 CAS 저장 -> {OUT_PATH}")
    print(f"분류표시 데이터 없는 물질(빈 시트): {empty_count}개")
    if errors:
        print(f"실패 {len(errors)}건:")
        for name, err in errors:
            print(f"  - {name}: {err}")


if __name__ == "__main__":
    sys.exit(main())
